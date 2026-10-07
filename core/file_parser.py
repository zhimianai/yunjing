import os
import io

from config import TEXT_EXTS, DOC_EXTS, UNSUPPORTED_EXTS


def looks_like_text(s: str) -> bool:
    if not s:
        return False
    text_chars = sum(1 for c in s if c.isprintable() or c in '\n\r\t')
    return text_chars / max(len(s), 1) > 0.90


def extract_text(filename: str, raw: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext in UNSUPPORTED_EXTS:
        supported = ', '.join(sorted(TEXT_EXTS))
        raise ValueError(
            f'不支持的文件格式：{ext or "无扩展名"}。\n'
            f'当前支持的文本/代码格式：{supported}\n'
            f'文档格式：PDF / .docx / .xlsx / .pptx\n'
            f'旧版 Office (.doc/.xls/.ppt) 和压缩包无法解析，请先转换格式。'
        )

    if ext in DOC_EXTS:
        return extract_doc(ext, raw)

    if not ext or ext in TEXT_EXTS:
        for enc in ('utf-8', 'utf-8-sig', 'gbk', 'gb18030', 'gb2312',
                    'big5', 'latin-1', 'cp936'):
            try:
                return raw.decode(enc)
            except (UnicodeDecodeError, LookupError):
                continue

    for enc in ('utf-8', 'gbk', 'gb18030', 'latin-1'):
        try:
            decoded = raw.decode(enc)
            if looks_like_text(decoded):
                return decoded
        except (UnicodeDecodeError, LookupError):
            continue

    size_kb = len(raw) // 1024
    raise ValueError(
        f'无法识别该文件类型（扩展名: {ext or "无"}，大小: {size_kb}KB）。\n'
        f'当前支持：文本/代码文件、PDF、.docx、.xlsx、.pptx\n'
        f'如果是图片、视频、压缩包或旧版 Office (.doc/.xls/.ppt)，请先转换格式再上传。'
    )


def extract_doc(ext: str, raw: bytes) -> str:
    _io = io

    try:
        if ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(_io.BytesIO(raw))
            parts = []
            for i, page in enumerate(reader.pages):
                try:
                    text = page.extract_text() or ''
                    parts.append(f"--- 第 {i+1} 页 ---\n{text}")
                except Exception:
                    parts.append(f"--- 第 {i+1} 页 ---\n[该页无法提取]")
            content = "\n\n".join(parts)
            if not content.strip():
                return "PDF 文件已识别，但未提取到文本内容（可能是扫描件或图片型 PDF）。"
            return content

        if ext == '.docx':
            from docx import Document
            doc = Document(_io.BytesIO(raw))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text)
            for table in doc.tables:
                parts.append("\n[表格]")
                for row in table.rows:
                    cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
                    parts.append(" | ".join(cells))
            return "\n".join(parts) if parts else "Word 文件已识别，但未提取到内容。"

        if ext in ('.xlsx', '.xlsm'):
            from openpyxl import load_workbook
            wb = load_workbook(_io.BytesIO(raw), data_only=True, read_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"=== 工作表: {ws.title} ===")
                row_count = 0
                for row in ws.iter_rows(values_only=True):
                    cells = [str(v) if v is not None else '' for v in row]
                    line = " | ".join(cells).strip()
                    if line:
                        parts.append(line)
                    row_count += 1
                    if row_count > 5000:
                        parts.append("...（表格过长，已截断）")
                        break
                parts.append("")
            return "\n".join(parts)

        if ext == '.pptx':
            from pptx import Presentation
            prs = Presentation(_io.BytesIO(raw))
            parts = []
            for i, slide in enumerate(prs.slides):
                slide_parts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_parts.append(shape.text.strip())
                    if shape.has_table:
                        for row in shape.table.rows:
                            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
                            slide_parts.append(" | ".join(cells))
                if slide_parts:
                    parts.append(f"--- 幻灯片 {i+1} ---")
                    parts.extend(slide_parts)
                    parts.append("")
            return "\n".join(parts) if parts else "PPT 文件已识别，但未提取到内容。"

    except ImportError:
        return f"缺少依赖库，无法解析 {ext}。请运行：pip install pypdf python-docx openpyxl python-pptx"
    except Exception as e:
        return f"解析 {ext} 文件失败：{str(e)}"

    return f"暂不支持的文档格式：{ext}"